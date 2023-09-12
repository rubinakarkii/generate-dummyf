from validator import Validation
from helper_functions import *
from datetime import datetime
import csv
import os
import json
import openpyxl
from docx import Document
from docx.shared import Pt

def prepare_csv(file_type,file_size,column_description): 
    data_type_value = {"str": generate_random_string(), "int": generate_random_int(), "datetime": datetime.now(), "float": generate_random_float(), "boolean": generate_random_boolean()}
    column_headers = []
    row = []
    for k,v in column_description.items():
        column_headers.append(k)
        row.append(data_type_value[v])

    new_file_path = get_path_to_create_new_file(file_type)
    with open(new_file_path, 'w', newline='') as csv_file:
        csv_writer = csv.writer(csv_file)
        current_file_size = 0
        csv_writer.writerow(column_headers)
        while(current_file_size<=file_size): 
            csv_writer.writerow(row)
            current_file_size = csv_file.tell() 
    
def prepare_json(file_type, file_size, column_description):
    data_type_value = {"str": generate_random_string(), "int": generate_random_int(), "datetime": datetime.now(), "float": generate_random_float(), "boolean": generate_random_boolean()}
    row = {k : data_type_value[v] for k,v in column_description.items()}
    new_file_path = get_path_to_create_new_file(file_type)
    data = []
    current_file_size = 0
    while(current_file_size<=file_size):
        with open(new_file_path, 'w') as json_file:            
            data.append(row)
            json.dump(data, json_file, default=str, indent=4)
            current_file_size = json_file.tell() 

def prepare_excel(file_type,file_size, column_description):
    data_type_value = {"str": generate_random_string(), "int": generate_random_int(), "datetime": datetime.now(), "float": generate_random_float(), "boolean": generate_random_boolean()}
    column_headers = []
    row = []
    for k,v in column_description.items():
        column_headers.append(k)
        row.append(data_type_value[v])

    new_file_path = get_path_to_create_new_file(file_type)
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    current_file_size = 0
    worksheet.append(column_headers)
    while(current_file_size<=file_size): 
        workbook.save(new_file_path)
        current_file_size = os.path.getsize(new_file_path)
        worksheet.append(row)
    workbook.close()

def prepare_txt(file_type,file_size):
    new_file_path = get_path_to_create_new_file(file_type)
    with open(new_file_path, 'a') as txt_file:
        current_file_size = 0
        while(current_file_size<=file_size): 
            txt = generate_random_string()
            txt_file.write(txt+ '\n')
            current_file_size = txt_file.tell() 

def prepare_docx(file_type,file_size,font_size):
    new_file_path = get_path_to_create_new_file(file_type)
    doc = Document()
    current_file_size = 0
    while(current_file_size<=file_size): 
        word_line = generate_random_string()
        paragraph = doc.add_paragraph(word_line)
        font = paragraph.runs[0].font
        font.size = Pt(font_size)
        doc.save(new_file_path)
        current_file_size = os.path.getsize(new_file_path)

def main(file_type, **kwargs):
    try:
        obj = Validation(**kwargs)
        obj.validate_args()
        obj.validate_file_size_limit()
        if file_type in ["csv","json","xlsx"]: obj.validate_column_description()
        if file_type == "docx": obj.validate_font_attributes()

        print(f"Generating dummy {file_type} file of size {kwargs['file_size']} bytes")
        if kwargs["file_size"]==0: zero_byte_file(file_type)
        else:
            if file_type=="csv": prepare_csv(file_type, kwargs["file_size"], kwargs["column_description"])
            if file_type=="json": prepare_json(file_type,kwargs["file_size"], kwargs["column_description"])
            if file_type=="xlsx": prepare_excel(file_type,kwargs["file_size"], kwargs["column_description"])
            if file_type=="txt": prepare_txt(file_type,kwargs["file_size"])
            if file_type=="docx": prepare_docx(file_type,kwargs["file_size"], kwargs["font_size"])
        print(f"Generated dummy {file_type} file in your downloads folder")
    except Exception as e:
        print(e)