import random 
from datetime import datetime
import csv
import os
import json
import openpyxl
import PyPDF2
from docx import Document

def validate_args(**kwargs):
    data_type_dict = {"file_size": int, "column_description" : dict}
    for key, value in kwargs.items():
        if not data_type_dict[key] == type(value):
                raise Exception(f"Error: {key} argument's data type should be {data_type_dict[key]}")

def get_path_to_create_new_file():
    downloads_folder = os.path.expanduser("~") + "/Downloads/"
    return downloads_folder

def zero_byte_file(file_type):
    with open(f'{get_path_to_create_new_file()}Dummy.{file_type}',"w"):
        return None

def generate_random_string():
    string = "abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ"
    return ''.join(random.sample(string,len(string)))

def generate_random_int():
    return random.getrandbits(32)

def generate_random_float():
    return random.random()

def generate_random_boolean():
    return random.choice([True, False])

def prepare_csv(file_size,column_description): 
    data_type_value = {"str": generate_random_string(), "int": generate_random_int(), "datetime": datetime.now(), "float": generate_random_float(), "boolean": generate_random_boolean()}
    column_headers = []
    row = []
    for k,v in column_description.items():
        column_headers.append(k)
        row.append(data_type_value[v])

    new_file_path = f'{get_path_to_create_new_file()}Dummy.csv'
    with open(new_file_path, 'w', newline='') as csv_file:
        csv_writer = csv.writer(csv_file)
        current_file_size = 0
        csv_writer.writerow(column_headers)
        while(current_file_size<=file_size): 
            csv_writer.writerow(row)
            current_file_size = csv_file.tell() 
    
def prepare_json(file_size, column_description):
    data_type_value = {"str": generate_random_string(), "int": generate_random_int(), "datetime": datetime.now(), "float": generate_random_float(), "boolean": generate_random_boolean()}
    row = {k : data_type_value[v] for k,v in column_description.items()}
    new_file_path = f'{get_path_to_create_new_file()}Dummy.json'
    data = []
    current_file_size = 0
    while(current_file_size<=file_size):
        with open(new_file_path, 'w') as json_file:            
            data.append(row)
            json.dump(data, json_file, default=str, indent=4)
            current_file_size = json_file.tell() 

def prepare_excel(file_size, column_description):
    data_type_value = {"str": generate_random_string(), "int": generate_random_int(), "datetime": datetime.now(), "float": generate_random_float(), "boolean": generate_random_boolean()}
    column_headers = []
    row = []
    for k,v in column_description.items():
        column_headers.append(k)
        row.append(data_type_value[v])

    new_file_path = f'{get_path_to_create_new_file()}Dummy.xlsx'
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    current_file_size = 0
    worksheet.append(column_headers)
    while(current_file_size<=file_size): 
        workbook.save(new_file_path)
        current_file_size = os.path.getsize(new_file_path)
        worksheet.append(row)
    workbook.close()

def prepare_txt(file_size):
    new_file_path = f'{get_path_to_create_new_file()}Dummy.txt'
    with open(new_file_path, 'a') as txt_file:
        current_file_size = 0
        while(current_file_size<=file_size): 
            txt = generate_random_string()
            txt_file.write(txt+ '\n')
            current_file_size = txt_file.tell() 

def prepare_pdf(file_size):
    new_file_path = f'{get_path_to_create_new_file()}Dummy.docx'
    # import ipdb;ipdb.set_trace()
    pdf_writer = PyPDF2.PdfFileWriter()
    current_file_size = 0
    while(current_file_size<=file_size): 
        pdf_line = generate_random_string()
        # current_file_size = .tell() 

def prepare_docx(file_size):
    new_file_path = f'{get_path_to_create_new_file()}Dummy.docx'
    doc = Document()
    current_file_size = 0
    while(current_file_size<=file_size): 
        word_line = generate_random_string()
        doc.add_paragraph(word_line)
        doc.save(new_file_path)
        current_file_size = os.path.getsize(new_file_path)

def main(file_type, **kwargs):
    try:
        validate_args(**kwargs)
        print(f"Generating dummy {file_type} file of size {kwargs['file_size']} bytes")
        if kwargs["file_size"]==0: zero_byte_file(file_type)
        else:
            if file_type=="csv": prepare_csv(kwargs["file_size"], kwargs["column_description"])
            if file_type=="json": prepare_json(kwargs["file_size"], kwargs["column_description"])
            if file_type=="xlsx": prepare_excel(kwargs["file_size"], kwargs["column_description"])
            if file_type=="txt": prepare_txt(kwargs["file_size"])
            if file_type=="pdf": prepare_pdf(kwargs["file_size"])
            if file_type=="docx": prepare_docx(kwargs["file_size"])
        print(f"Generated dummy {file_type} file in your downloads folder")
    except Exception as e:
        print(e)